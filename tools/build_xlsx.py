import sys
sys.path.insert(0, '/home/claude/web/tools')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter as CL
from lm01_data import *

F = 'Arial'; wb = Workbook()
def sh(name, hdr, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet(name); ws.title = name
    ws.append(hdr)
    for r in rows: ws.append(list(r))
    for c in ws[1]:
        c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', start_color='1F3864')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths, 1): ws.column_dimensions[CL(i)].width = w
    ws.freeze_panes = 'B2'
    t = Table(displayName='t_' + name, ref=f'A1:{CL(len(hdr))}{len(rows)+1}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True); ws.add_table(t)
    return ws

sh('Module', list(MODULE.keys()), [list(MODULE.values())], [12,8,6,10,40,10,26,26,10,10,10,8,12], True)

lid = lambda l: f'{M}-{l.replace(".", "")}'
SRCS = ['Curriculum', 'Schweser QBank', 'Schweser Quiz', 'UWorld', 'AnalystPrep']
rows = []
for i, x in enumerate(LOS, 1):
    r = i + 1
    qf = [f'=COUNTIFS(Questions!$D:$D,$A{r},Questions!$B:$B,"{s}",Questions!$F:$F,"In LM1")' for s in SRCS]
    qf[0] = f'=COUNTIFS(Questions!$D:$D,$A{r},Questions!$B:$B,"Curriculum*",Questions!$F:$F,"In LM1")'
    cf = f'=COUNTIFS(Conflicts!$B:$B,$A{r})'
    xm = f'=COUNTIFS(Extras!$B:$B,$A{r},Extras!$E:$E,"Must")'
    rows.append([lid(x[0]), M, x[0], i, x[1], x[2], x[3], x[4], x[5], x[6], x[7], x[8], x[9], x[10], x[11], x[10]-2, x[11]-2, x[12],
                 *qf, f'=SUM(S{r}:W{r})', cf, xm, None, None])
sh('LOS', ['los_id','module_id','los_code','order','los_text','recommendation','schweser_coverage_pct','curriculum_coverage_pct','why_recommended','why_not_other',
           'sch_page_from','sch_page_to','sch_anchor','curr_printed_from','curr_printed_to','curr_pdf_from','curr_pdf_to','curr_anchor',
           'q_curriculum','q_schweser_qbank','q_schweser_quiz','q_uworld','q_analystprep','q_total','n_conflicts','n_must_extras','user_preference','user_note'],
   rows, [16,10,7,6,45,10,9,9,50,40,7,7,30,7,7,7,7,30,8,8,8,8,8,8,8,8,10,25])

xr = []; cnt = {}
for los, d, t, p, txt, pg, anc, lk in EXTRAS:
    cnt[los] = cnt.get(los, 0) + 1
    xr.append([f'{lid(los)}-X{cnt[los]:02d}', lid(los), cnt[los], d, p, t, txt,
               'Schweser' if d == 'SCH' else 'Curriculum', pg, pg - 2 if d == 'CUR' else pg, anc, lk])
sh('Extras', ['extra_id','los_id','order','direction','priority','title','full_text','page_source','printed_page','pdf_page','anchor','exam_link'],
   xr, [22,16,6,9,8,32,80,11,8,8,28,30])

sh('Conflicts', ['conflict_id','los_id','topic','schweser_says','sch_page','curriculum_says','curr_printed_page','use_for_exam','exam_risk','qbank_evidence'],
   [[c[0], lid(c[1])] + list(c[2:]) for c in CONFLICTS], [9,16,26,45,7,50,8,40,9,35])

code = {'Curriculum':'CUR','Curriculum KC':'CKC','Schweser QBank':'SQB','Schweser Quiz':'SQZ','UWorld':'UW','AnalystPrep':'AP'}
sh('Questions', ['question_ref','source','q_no','los_id','answer','scope','quality_flag','mentor_note'],
   [[f'{M}-{code[s]}-{n:04d}', s, n, lid(l), a, sc, fl, note] for s, n, l, a, sc, fl, note in Q], [22,16,7,16,8,18,12,60])

sh('Sources', ['source_id','name','role','status','note'], SOURCES, [10,30,14,24,60])

ws = wb.create_sheet('README'); ws.column_dimensions['A'].width = 120
for i, t in enumerate([
 'FSA-LM01 — Study Map data file (template v1.1)', '',
 'Every sheet is an Excel Table; one row = one record with a stable ID. No merged cells. Counts on the LOS sheet are formulas over Questions/Extras/Conflicts.',
 'recommendation: S = Schweser only · C = Curriculum instead of Schweser · S+X = Schweser + Extras of this LOS.',
 'Extras direction: CUR = the Curriculum adds this (read it if you study from Schweser). SCH = Schweser adds this (read it if you study from the Curriculum).',
 'priority: Must = examinable, full text · Useful = context, full text · Example = short summary · Skip = title only.',
 'Rule: Schweser + CUR Extras = Curriculum + SCH Extras = full exam coverage.',
 'Pages: sch_* = page in the cut Schweser file. curr_printed = number printed in the book; curr_pdf = page in the cut Curriculum file (printed − 2).',
 'Questions: index only (no question text). scope = In LM1 / Later module / Off-curriculum detail. Only "In LM1" is counted on the LOS sheet.',
 'Question text is imported by the website from your own PDFs, on your device.',
 'user_preference and user_note are filled by you (or by the website).'], 1):
    ws.cell(i, 1, t).font = Font(name=F, bold=(i == 1), size=12 if i == 1 else 10)
out = sys.argv[1]; wb.save(out); print(out)

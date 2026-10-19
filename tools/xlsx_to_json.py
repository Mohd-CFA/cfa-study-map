"""Convert a Study Map workbook (after recalculation) to the JSON the website reads.
Usage: python xlsx_to_json.py FSA-LM01_StudyMap_v1.1.xlsx data/FSA-LM01.json"""
import sys, json
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1], data_only=True)
def rows(name):
    ws = wb[name]; it = ws.iter_rows(values_only=True); h = next(it)
    return [dict(zip(h, r)) for r in it if any(v is not None for v in r)]
d = {'module': rows('Module')[0]}
for s in ['LOS', 'Extras', 'Conflicts', 'Questions', 'Sources']: d[s.lower()] = rows(s)
json.dump(d, open(sys.argv[2], 'w'), ensure_ascii=False, indent=1, default=str)
print(sys.argv[2], {k: len(v) for k, v in d.items() if isinstance(v, list)})
